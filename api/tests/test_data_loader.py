import hashlib
from io import BytesIO

import pytest

from src.data_loader import load_excel, locate_excel, number
from src.signals import product_rows, sales_signal


def test_headers_and_all_sheets(bundle):
    assert len(bundle.tables) == 11
    assert all(v == 3 for v in bundle.headers.values())
    assert bundle.reference.isoformat() == "2026-09-09"
    assert len(bundle.tables["Vendas_12m"]) == 1440
    assert len(bundle.tables["Regras_Negocio"]) == 12


def test_full_catalog_all_channels(bundle):
    rows = product_rows(bundle)
    assert len(rows) == 120
    assert len({r["sku"] for r in rows}) == 40
    assert {r["canal"] for r in rows} == {"Loja física", "E-commerce", "Marketplace"}
    assert all(isinstance(r["sku"],str) for r in rows)


def test_original_preserved(workbook):
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    bundle = load_excel(workbook)
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before == bundle.source_hash


def test_currency_percent_conversion():
    assert number("R$ 1.234,56") == 1234.56
    assert number("25%", True) == .25
    assert number(25,True) == .25
    assert number(.25,True) == .25
    assert number("inválido") is None


def test_bad_upload_rejected():
    with pytest.raises(ValueError, match="xlsx"):
        load_excel(b"abc", "base.exe")
    with pytest.raises(ValueError, match="inválido"):
        load_excel(b"abc", "base.xlsx")


def test_monthly_sale_not_exact(bundle):
    signal = sales_signal(bundle,"PP-0001","Loja física",bundle.reference)
    assert signal["origem"] == "estimado"
    assert not signal["comprovado"]
    assert signal["ultima_venda"] is None
    assert signal["ultimo_periodo"] == "2026-08"


def test_import_roundtrip(service):
    bundle = service.bundle()
    assert len(bundle.tables["Precos_Atuais"]) == 120
    recs = service.analyze()
    assert len(recs) == 120
    assert all("reference" in r for r in recs)


def test_missing_price_kept_for_review(service, monkeypatch):
    bundle = service.bundle()
    bundle.tables["Precos_Atuais"] = bundle.tables["Precos_Atuais"].iloc[1:].copy()
    monkeypatch.setattr(service, "bundle", lambda con=None:bundle)
    rows = service.analyze()
    assert len(rows) == 120
    assert any(r["preco_sugerido"] is None and "ausentes" in r["justificativa"] for r in rows)


def test_duplicate_key_rejected_without_changing_original(workbook):
    import openpyxl
    book = openpyxl.load_workbook(workbook)
    book["Precos_Atuais"].cell(5,2).value = book["Precos_Atuais"].cell(4,2).value
    copy = BytesIO()
    book.save(copy)
    with pytest.raises(ValueError,match="duplicadas"):
        load_excel(copy.getvalue(),"duplicada.xlsx")


def test_header_detects_extra_title_rows(workbook):
    import openpyxl
    book = openpyxl.load_workbook(workbook)
    book["Produtos"].insert_rows(1,2)
    copy = BytesIO()
    book.save(copy)
    result = load_excel(copy.getvalue(),"cabecalho.xlsx")
    assert result.headers["Produtos"] == 5
    assert len(result.tables["Produtos"]) == 40


def test_upload_missing_sheet_rejected(workbook):
    import openpyxl
    book = openpyxl.load_workbook(workbook)
    del book["Estoque"]
    copy = BytesIO()
    book.save(copy)
    with pytest.raises(ValueError,match="Abas ausentes"):
        load_excel(copy.getvalue(),"incompleta.xlsx")
