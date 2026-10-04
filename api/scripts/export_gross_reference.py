"""Gera a referência do motor do Allan com a margem bruta sobre reposição.

Truque: zera impostos, frete e taxa do canal e iguala o custo total à reposição.
Com isso, a margem de contribuição do motor vira (preço − reposição) ÷ preço.
Serve para conferir a prévia do front (web/scripts/teste-recalculo.mts) no modo "gross".

Uso: python scripts/export_gross_reference.py
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from src.database import Database
from src.recommendation_service import RecommendationService

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "web" / "scripts" / "gross-reference.json"
HEADER_ROW = 3


def gross_workbook(source: Path, target: Path) -> None:
    shutil.copy(source, target)
    wb = load_workbook(target)
    ws = wb["Custos_Margens"]
    cols = {str(c.value).strip(): c.column for c in ws[HEADER_ROW] if c.value}
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        replacement = ws.cell(row, cols["Custo reposição"]).value
        if replacement is None:
            continue
        ws.cell(row, cols["Impostos estimados"]).value = 0
        ws.cell(row, cols["Frete rateado"]).value = 0
        ws.cell(row, cols["Taxa canal"]).value = 0
        ws.cell(row, cols["Custo total referência"]).value = replacement
    wb.save(target)


def run(source: Path) -> dict:
    tmp = Path(tempfile.mkdtemp())
    workbook = tmp / "gross.xlsx"
    gross_workbook(source, workbook)
    service = RecommendationService(Database(os.path.join(tmp, "x.db")))
    service.import_excel(str(workbook), user="referencia")
    service.generate("referencia")
    return {
        f"{r['sku']}|{r['canal']}": {
            "acao": r["acao"], "risco": r["nivel_risco"],
            "preco_sugerido": float(r["preco_sugerido"]) if r["preco_sugerido"] is not None else None,
            "preco_minimo": float(r["preco_minimo"]) if r["preco_minimo"] is not None else None,
        }
        for r in service.latest()
    }


if __name__ == "__main__":
    payload = {
        "oficial": run(ROOT / "data" / "base-oficial.xlsx"),
        "sintetico": run(ROOT / "data" / "cenarios-sinteticos.xlsx"),
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=1))
    print({k: len(v) for k, v in payload.items()})
