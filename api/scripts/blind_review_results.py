"""Resultados da revisão cega da Semana 4: respostas de Allan e Caio contra o gabarito do app.

As respostas do Allan vêm da planilha dele no Drive (Grupo 12 > Semana 4, salva em 05/10/2026 12:10), lidas pelo conector.
As do Caio vêm da planilha baixada (as fórmulas =H2 etc. são lidas pelo valor calculado).
"""
from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import load_workbook

OUT = Path.home() / "Documents/popular-pet/[04] outputs/semana-4/data"
OUT.mkdir(parents=True, exist_ok=True)
KEY = Path.home() / "Documents/popular-pet/[04] outputs/semana-3/revisao-cega/GABARITO - nao compartilhar.xlsx"

ALLAN = {1: ("BAIXAR", 149.50), 2: ("MANTER", None), 3: ("BAIXAR", 145.41), 4: ("MANTER", None), 5: ("MANTER", None), 6: ("MANTER", None),
         7: ("BAIXAR", 36.76), 8: ("BAIXAR", 17.47), 9: ("BAIXAR", 79.43), 10: ("MANTER", None), 11: ("MANTER", None), 12: ("MANTER", None)}


def caio_answers(path: str) -> dict[int, tuple]:
    ws = load_workbook(path, data_only=True)["Combinações"]
    return {r[0]: (r[20], r[21]) for r in ws.iter_rows(min_row=2, max_row=13, values_only=True)}


caio = caio_answers(os.path.expanduser("~/Downloads/Revisao cega - Caio.xlsx"))
key = {r[0]: {"app_contrib": r[7], "app_gross": r[9]} for r in load_workbook(KEY)["Gabarito do app"].iter_rows(min_row=2, values_only=True)}


def agree(answers: dict, k: str) -> int:
    return sum(1 for n in answers if answers[n][0] == key[n][k])


result = {
    "agree_contrib": {"allan": agree(ALLAN, "app_contrib"), "caio": agree(caio, "app_contrib")},
    "agree_gross": {"allan": agree(ALLAN, "app_gross"), "caio": agree(caio, "app_gross")},
    "allan_vs_caio": sum(1 for n in ALLAN if ALLAN[n][0] == caio[n][0]),
    "raise_answers": sum(1 for h in (ALLAN, caio) for n in h if h[n][0] == "SUBIR"),
}
(OUT / "results.json").write_text(json.dumps(result, ensure_ascii=False, indent=1))
print(result)
assert result["agree_contrib"] == {"allan": 0, "caio": 1}, result
assert result["agree_gross"] == {"allan": 4, "caio": 5}, result
assert result["allan_vs_caio"] == 5 and result["raise_answers"] == 0, result
