"""Resultados da rodada 2 da revisão cega (Semana 5): respostas de Allan e Caio contra as previsões congeladas do app.

Uso: PYTHONPATH=. python scripts/round2_results.py <planilha do Allan.xlsx> <planilha do Caio.xlsx>
Sem argumentos, grava {"status": "pendente"}: o dossiê trata o caso como "rodada 2 não concluída".
Antes de ler as respostas, confere o hash das previsões gravado em docs/round2-predictions.sha256.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HOME = Path.home() / "Documents/popular-pet"
DATA = HOME / "[04] outputs/semana-5/data"
APP = HOME / "[03] projects/pricing-intelligence/app"
OUT = DATA / "results2.json"

if len(sys.argv) != 3:
    OUT.write_text(json.dumps({"status": "pendente"}, ensure_ascii=False, indent=1))
    print("sem planilhas: status pendente")
    raise SystemExit(0)

text = (DATA / "round2-predictions.json").read_text()
expected = (APP / "docs/round2-predictions.sha256").read_text().split()[0]
assert hashlib.sha256(text.encode()).hexdigest() == expected, "as previsões mudaram depois do congelamento"
pred = {c["n"]: c for c in json.loads(text)["casos"]}
base = {r["id"]: r for r in json.loads((APP / "web/src/data/oficial.json").read_text())["recomendacoes"]}
ids = {c["n"]: c["id"] for c in pred.values()}


def answers(path: str) -> dict[int, dict]:
    ws = load_workbook(path, data_only=True)["Combinações"]
    out = {}
    for r in ws.iter_rows(min_row=2, max_row=13, values_only=True):
        out[r[0]] = {"acao": r[20], "preco": r[21], "formula": r[22], "tempo": r[24]}
    missing = [n for n, a in out.items() if not a["acao"]]
    assert not missing, f"combinações sem ação em {path}: {missing}"
    return out


def stats(a: dict[int, dict]) -> dict:
    band = r02 = below_gross = 0
    times = []
    for n, x in a.items():
        r = base[ids[n]]
        price = x["preco"] if x["acao"] in ("SUBIR", "BAIXAR") and x["preco"] else r["preco_atual"]
        if abs(price / r["preco_mercado"] - 1) <= 0.02:
            band += 1
        if abs(price / r["preco_atual"] - 1) > 0.05:
            r02 += 1
        if (price - r["custo"]["reposicao"]) / price < r["margem"]["minima"] - 1e-9:
            below_gross += 1
        if isinstance(x["tempo"], (int, float)):
            times.append(float(x["tempo"]))
    return {
        "acordo_bruta": sum(1 for n, x in a.items() if x["acao"] == pred[n]["acao"]),
        "acordo_contribuicao": sum(1 for n, x in a.items() if x["acao"] == pred[n]["contribuicao"]["acao"]),
        "subir_escolhidos": sum(1 for x in a.values() if x["acao"] == "SUBIR"),
        "na_faixa": band, "passa_5pct": r02, "abaixo_minimo_bruta": below_gross,
        "tempo_total_min": sum(times) if len(times) == 12 else None,
        "tempo_medio_min": sum(times) / len(times) if len(times) == 12 else None,
    }


allan, caio = answers(sys.argv[1]), answers(sys.argv[2])
res = {
    "status": "concluida", "hash_previsoes": expected,
    "allan": stats(allan), "caio": stats(caio),
    "allan_vs_caio": sum(1 for n in allan if allan[n]["acao"] == caio[n]["acao"]),
    "previsao_app": {a: sum(1 for c in pred.values() if c["acao"] == a) for a in ("SUBIR", "BAIXAR", "MANTER", "REVISAR")},
}
OUT.write_text(json.dumps(res, ensure_ascii=False, indent=1))
print(json.dumps(res, ensure_ascii=False, indent=1))
