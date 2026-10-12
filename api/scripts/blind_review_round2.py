"""Rodada 2 da revisão cega (Semana 5): 12 combinações novas, sorteio com semente fixa e previsões do app congeladas.

Regras do sorteio: nenhum produto nem combinação da rodada 1; 3 combinações de cada ação do app (padrão: margem bruta com teto de mercado);
4 por canal; as 3 curvas e pelo menos 5 categorias. As previsões são gravadas, com hash, antes de qualquer resposta.
Uso: PYTHONPATH=. python scripts/blind_review_round2.py
"""
from __future__ import annotations

import collections
import datetime as dt
import hashlib
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HOME = Path.home() / "Documents/popular-pet"
DATA = HOME / "[04] outputs/semana-5/data"
OUT = HOME / "[04] outputs/semana-5/revisao-cega-2"
APP = HOME / "[03] projects/pricing-intelligence/app"
base = json.loads((APP / "web/src/data/oficial.json").read_text())
outputs = json.loads((DATA / "app-outputs.json").read_text())
round1 = set(json.loads((HOME / "[04] outputs/semana-3/revisao-cega/selecao.json").read_text()))
used_skus = {i.split("|")[0] for i in round1}
recs = {r["id"]: r for r in base["recomendacoes"]}
pool = [i for i in recs if i not in round1 and i.split("|")[0] not in used_skus]
by_action = collections.defaultdict(list)
for i in pool:
    by_action[outputs[i]["gross"]["acao"]].append(i)
ACTIONS = ["SUBIR", "BAIXAR", "MANTER", "REVISAR"]
assert all(len(by_action[a]) >= 3 for a in ACTIONS), {a: len(by_action[a]) for a in ACTIONS}

best = None
for seed in range(300000):
    rnd = random.Random(seed)
    pick = [i for a in ACTIONS for i in rnd.sample(by_action[a], 3)]
    if len({i.split("|")[0] for i in pick}) < 12:
        continue
    ch = collections.Counter(recs[i]["canal"] for i in pick)
    if set(ch.values()) != {4}:
        continue
    cur = collections.Counter(recs[i]["curva"] for i in pick)
    if len(cur) < 3 or min(cur.values()) < 3:
        continue
    cats = len({recs[i]["categoria"] for i in pick})
    if cats < 5:
        continue
    stale = sum(recs[i]["concorrencia"]["desatualizado"] for i in pick)
    score = cats * 10 + stale * 3
    if best is None or score > best[0]:
        best = (score, seed, pick)
score, seed, pick = best
random.Random(11).shuffle(pick)  # ordem de apresentação sem agrupar por ação
print("semente", seed, "| canais", collections.Counter(recs[i]["canal"] for i in pick), "| curvas", collections.Counter(recs[i]["curva"] for i in pick),
      "| categorias", len({recs[i]["categoria"] for i in pick}), "| desatualizados", sum(recs[i]["concorrencia"]["desatualizado"] for i in pick))

# ---- previsões congeladas
pred = {"gerado_em": dt.datetime.now().isoformat(timespec="seconds"), "semente": seed, "app": "pet-pricing.vercel.app, commit 54d423d, margem bruta com teto de mercado",
        "casos": [{"n": n, "id": i, **outputs[i]["gross"], "contribuicao": outputs[i]["contribution"]} for n, i in enumerate(pick, 1)]}
text = json.dumps(pred, ensure_ascii=False, indent=1)
(DATA / "round2-predictions.json").write_text(text)
digest = hashlib.sha256(text.encode()).hexdigest()
(DATA / "round2-predictions.sha256").write_text(f"{digest}  round2-predictions.json\n")
print("sha256", digest)

# ---- planilhas
ref = dt.date(2026, 9, 9)

def promo(r):
    for p in r["promocoes"]:
        a = dt.datetime.strptime(p["inicio"], "%d/%m/%Y").date(); b = dt.datetime.strptime(p["fim"], "%d/%m/%Y").date()
        if a <= ref <= b:
            return f'Sim: {p["campanha"]} ({p["inicio"][:5]} a {p["fim"][:5]}, {round(p["desconto"] * 100)}% de desconto)'
    return "Não"

HEAD = ["Nº", "SKU", "Produto", "Categoria", "Curva", "Canal", "Preço atual (R$)", "Concorrente A (R$)", "Concorrente B (R$)", "Concorrente C (R$)",
        "Mediana do mercado (R$)", "Custo de reposição (R$)", "Impostos (R$)", "Frete rateado (R$)", "Taxa do canal", "Margem mínima do SKU/canal",
        "Cobertura de estoque (dias)", "Vendas nos últimos 3 meses (un.)", "Promoção ativa em 09/09/2026?", "Coleta de concorrência"]
ANSWER = ["SUA AÇÃO (SUBIR, BAIXAR, MANTER ou REVISAR)", "SEU PREÇO (R$), se for subir ou baixar", "Fórmula de margem usada (bruta ou contribuição)",
          "Por quê (uma linha)", "Tempo gasto nesta combinação (minutos)"]

def rows():
    out = []
    for n, i in enumerate(pick, 1):
        r = recs[i]
        comp = [c["preco"] if c["disponivel"] == "Sim" else f'{c["preco"]:.2f} (indisponível)'.replace(".", ",") for c in r["concorrentes"]]
        comp += ["sem coleta"] * (3 - len(comp))
        out.append([n, r["sku"], r["produto"], r["categoria"], r["curva"], r["canal"], r["preco_atual"], comp[0], comp[1], comp[2], r["preco_mercado"],
                    r["custo"]["reposicao"], r["custo"]["impostos"], r["custo"]["frete"], r["custo"]["taxa_canal"], r["margem"]["minima"],
                    r["estoque"]["cobertura_dias"], sum(m["quantidade"] for m in r["vendas"][-3:]), promo(r),
                    "Desatualizada (mais de 48 h)" if r["concorrencia"]["desatualizado"] else "Em dia"])
    return out

INSTR = [
    ("Revisão cega, rodada 2: o que fazer", True),
    ("Objetivo: comparar o que você decidiria com o que o Pet Pricing recomenda, em 12 combinações novas que ninguém usou para ajustar o app. As respostas do app já estão guardadas, e ninguém as viu.", False),
    ("", False),
    ("Como responder (cerca de 30 minutos)", True),
    ("1. Não abra o app nem as respostas de outra pessoa até quinta (15/10).", False),
    ("2. Cronometre: anote, em cada linha, quantos minutos você levou. Esse é o tempo manual que vamos comparar com o do app.", False),
    ("3. Na aba \"Combinações\", há 12 pares produto × canal. Preencha as cinco colunas azuis do final.", False),
    ("4. Ação: SUBIR, BAIXAR, MANTER ou REVISAR. Use REVISAR quando uma pessoa precisaria olhar o caso antes de mexer no preço.", False),
    ("5. Preço: o valor em reais que você aplicaria. Deixe em branco se for MANTER ou REVISAR.", False),
    ("6. Fórmula de margem: diga qual usou. Bruta = (preço − custo de reposição) ÷ preço. Contribuição = (preço × (1 − taxa do canal) − reposição − impostos − frete) ÷ preço.", False),
    ("7. Por quê: uma linha, com as suas palavras.", False),
    ("8. Prazo: quarta (14/10) às 18h. Mande só um \"pronto\" no grupo, sem contar as respostas.", False),
    ("", False),
    ("Dados de apoio", True),
    ("Base oficial do desafio, referência em 09/09/2026. Preços e custos fictícios. A aba \"Regras do desafio\" traz as 12 regras.", False),
]
thin = Side(style="thin", color="D9D4E0")

def build(path, owner):
    wb = Workbook(); ws = wb.active; ws.title = "Leia primeiro"; ws.column_dimensions["A"].width = 120
    ws.append([f"Revisão cega, rodada 2: {owner}"]); ws["A1"].font = Font(bold=True, size=16, color="5A3580")
    for t, head in INSTR:
        ws.append([t]); c = ws.cell(ws.max_row, 1); c.alignment = Alignment(wrap_text=True, vertical="top")
        if head: c.font = Font(bold=True, size=12, color="5A3580")
    wc = wb.create_sheet("Combinações"); wc.append(HEAD + ANSWER)
    for row in rows(): wc.append(row)
    nh = len(HEAD)
    for col in range(1, nh + len(ANSWER) + 1):
        c = wc.cell(1, col); c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="764EA0" if col <= nh else "2F6FB3")
    for r in range(2, 14):
        for col in range(1, nh + len(ANSWER) + 1):
            c = wc.cell(r, col); c.border = Border(top=thin, bottom=thin, left=thin, right=thin); c.alignment = Alignment(wrap_text=True, vertical="top")
            if col > nh: c.fill = PatternFill("solid", fgColor="EAF2FB")
        for col in range(7, 15):
            if isinstance(wc.cell(r, col).value, float): wc.cell(r, col).number_format = "0.00"
        wc.cell(r, 15).number_format = "0.0%"; wc.cell(r, 16).number_format = "0%"; wc.cell(r, nh + 2).number_format = "0.00"
    wc.row_dimensions[1].height = 62
    for i, w in enumerate([5, 9, 30, 15, 7, 13, 11, 14, 14, 14, 12, 12, 10, 10, 9, 11, 11, 13, 34, 20, 24, 18, 20, 40, 16], 1):
        wc.column_dimensions[wc.cell(1, i).column_letter].width = w
    wc.freeze_panes = "D2"
    dv = DataValidation(type="list", formula1='"SUBIR,BAIXAR,MANTER,REVISAR"', allow_blank=True); wc.add_data_validation(dv); dv.add("U2:U13")
    dv2 = DataValidation(type="list", formula1='"Bruta,Contribuição"', allow_blank=True); wc.add_data_validation(dv2); dv2.add("W2:W13")
    wr = wb.create_sheet("Regras do desafio"); wr.append(["Regra", "Tema", "O que diz", "O que fazer"])
    for r in base["regras"]: wr.append([r["id"], r["tema"], r["regra"], r["acao"]])
    for c in wr[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="764EA0")
    for col, w in zip("ABCD", (8, 16, 90, 24)): wr.column_dimensions[col].width = w
    for row in wr.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)

for owner in ("Allan", "Caio"):
    build(OUT / f"Revisao cega 2 - {owner}.xlsx", owner)
# gabarito (fica só no computador da Safira)
wb = Workbook(); ws = wb.active; ws.title = "Gabarito do app (rodada 2)"
ws.append(["Nº", "Combinação", "Produto", "Canal", "Curva", "Preço atual", "Mediana", "Ação (bruta, padrão)", "Preço sugerido (bruta)", "Preço mínimo (bruta)", "Ação (contribuição)", "Preço sugerido (contribuição)"])
for n, i in enumerate(pick, 1):
    r, o = recs[i], outputs[i]
    ws.append([n, i, r["produto"], r["canal"], r["curva"], r["preco_atual"], r["preco_mercado"], o["gross"]["acao"], o["gross"]["preco_sugerido"], o["gross"]["preco_minimo"], o["contribution"]["acao"], o["contribution"]["preco_sugerido"]])
for c in ws[1]: c.font = Font(bold=True)
wb.save(OUT / "GABARITO 2 - nao compartilhar.xlsx")
(DATA / "round2-selection.json").write_text(json.dumps(pick, ensure_ascii=False, indent=1))
